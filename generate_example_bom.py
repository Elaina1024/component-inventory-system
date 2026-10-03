"""Generate a synthetic BOM for trying the application safely."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


DEMO_ROWS = [
    ("MCU-202405-000001", "示例主控芯片", "LQFP-48", 1, "A-01-01"),
    ("RES-202405-000001", "示例贴片电阻", "0805", 2, "B-02-05"),
    ("CAP-202405-000001", "示例贴片电容", "0805", 4, "C-01-01"),
]


def main():
    workbook = Workbook()
    workbook.properties.creator = "Inventory System Demo"
    workbook.properties.lastModifiedBy = "Inventory System Demo"
    workbook.properties.title = "Synthetic BOM example"
    workbook.properties.description = "Fictional records for software demonstration only"
    sheet = workbook.active
    sheet.title = "示例BOM"
    sheet.append(["内部料号", "名称", "封装", "用量", "库位"])
    for row in DEMO_ROWS:
        sheet.append(row)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="34699A")
        cell.alignment = Alignment(horizontal="center")
    for column, width in {"A": 24, "B": 20, "C": 14, "D": 10, "E": 14}.items():
        sheet.column_dimensions[column].width = width

    target = Path(__file__).resolve().with_name("example_bom.xlsx")
    workbook.save(target)
    print(f"已生成匿名示例 BOM：{target.name}（{len(DEMO_ROWS)} 行）")


if __name__ == "__main__":
    main()
